import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, Optional

class ReportGenerator:
    """Generates professional Word reports from tagged budget data based on a template."""
    
    def __init__(self, output_dir: str = "reports"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(exist_ok=True)

    def generate_summary_report(
        self, 
        df: pd.DataFrame, 
        metadata: Dict[str, Any],
        metrics: Dict[str, Any]
    ) -> Path:
        """Create a Word summary report aligned with the official template."""
        doc = Document()
        
        # Style setup (optional: would be better to load a template .docx but we create from scratch)
        self._add_title(doc, f"CLIMATE BUDGET TAGGING REPORT - {metadata.get('country', 'N/A')}")
        
        # 1. Basic Report Data
        self._add_section_header(doc, "1. Basic Report Data")
        table = doc.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        data_rows = [
            ("Country / Entity", metadata.get("country", "N/A")),
            ("Administrative Level", metadata.get("admin_level", "National")),
            ("Budget Stage", metadata.get("budget_stage", "Execution")),
            ("Fiscal Period", f"{metadata.get('year', 'N/A')}"),
            ("Methodology", metadata.get("method", "Rio Markers")),
            ("Data Source", "Climate Budget Tagging Automation Tool (AI + Human Validation)"),
            ("Generation Date", datetime.now().strftime('%Y-%m-%d'))
        ]
        for label, value in data_rows:
            row_cells = table.add_row().cells
            row_cells[0].text = label
            row_cells[1].text = str(value)

        # 2. Executive Summary
        self._add_section_header(doc, "2. Executive Summary")
        p = doc.add_paragraph()
        p.add_run(f"Total climate-relevant expenditure: ").bold = True
        p.add_run(f"{metrics.get('green_pct', 0):.1f}% of total analyzed budget.")
        
        p = doc.add_paragraph()
        p.add_run(f"Absolute Climate alignment: ").bold = True
        p.add_run(f"${metrics.get('green_amount', 0):,.2f} {metadata.get('currency', 'USD')}.")

        # Objective Split
        obj_split = metrics.get("objective_split", {})
        split_text = (
            f"Objective split — Mitigation: {obj_split.get('mitigation', 0):.1f}% | "
            f"Adaptation: {obj_split.get('adaptation', 0):.1f}% | "
            f"Cross-cutting: {obj_split.get('crosscutting', 0):.1f}%."
        )
        doc.add_paragraph(split_text)
        
        # Capital Share
        doc.add_paragraph(f"Capital share of climate expenditures: {metrics.get('capital_share', 0):.1f}%.")

        # 3. Climate Budget Trends
        self._add_section_header(doc, "3. Climate Budget Trends")
        trend = metrics.get("trend_pp", 0)
        trend_text = "stable"
        if trend > 0: trend_text = f"increased by {trend} percentage points"
        elif trend < 0: trend_text = f"decreased by {abs(trend)} percentage points"
        
        doc.add_paragraph(f"Trend Analysis: Overall climate relevance has {trend_text} compared to previous periods available in the dataset.")

        # 4. Sectoral Breakdown
        self._add_section_header(doc, "4. Sectoral Breakdown")
        sector_df = metrics.get("sectoral_breakdown")
        if sector_df is not None and not sector_df.empty:
            table = doc.add_table(rows=1, cols=3)
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Sector'
            hdr_cells[1].text = 'Share (%)'
            hdr_cells[2].text = f'Amount ({metadata.get("currency", "USD")})'
            
            for _, row in sector_df.iterrows():
                row_cells = table.add_row().cells
                row_cells[0].text = str(row['sector'])
                row_cells[1].text = f"{row['pct']:.1f}%"
                row_cells[2].text = f"{row['amount']:,.2f}"
        else:
            doc.add_paragraph("Sectoral data unavailable.")

        # 5. Quality, Validation & Oversight
        self._add_section_header(doc, "5. Quality, Validation & Oversight")
        conf = metrics.get("confidence_dist", {})
        total_items = sum(conf.values())
        doc.add_paragraph(f"Total budget items tagged: {total_items}")
        doc.add_paragraph(f"Confidence Distribution — High: {conf.get('HIGH', 0)} | Medium: {conf.get('MEDIUM', 0)} | Low: {conf.get('LOW', 0)}")

        # 6. Practitioner Appendix
        self._add_section_header(doc, "6. Practitioner Appendix - Top 20 Items")
        top_items = df.sort_values(by=metrics.get('amount_col', df.columns[0]), ascending=False).head(20)
        
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Program/Item'
        hdr_cells[1].text = 'Sector'
        hdr_cells[2].text = 'Amount'
        hdr_cells[3].text = 'Tag'
        hdr_cells[4].text = 'Confidence'
        
        for _, row in top_items.iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row.get('prog1', row.get('project_title', 'N/A')))[:50]
            row_cells[1].text = str(row.get('func1', 'N/A'))
            row_cells[2].text = f"{row.get(metrics.get('amount_col', 'amount'), 0):,.0f}"
            row_cells[3].text = str(row.get('climate_tag', 'N/A'))
            row_cells[4].text = str(row.get('confidence', 'N/A'))

        filename = f"Climate_Budget_Report_{metadata.get('country', 'Analysis')}_{metadata.get('year', 'FY')}.docx"
        file_path = self.output_dir / filename
        doc.save(file_path)
        return file_path

    def _add_title(self, doc, text):
        title = doc.add_heading(text, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _add_section_header(self, doc, text):
        heading = doc.add_heading(text, level=1)
        # Optional: styling
        # run = heading.runs[0]
        # run.font.color.rgb = RGBColor(0x1B, 0x4F, 0x72) # Dark Blue
